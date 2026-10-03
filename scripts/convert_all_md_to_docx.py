"""Convert all Markdown (.md) files in the repository to Word (.docx) files.

Saves all converted .docx files into the 'word/' directory.
Uses python-docx to properly style headings, lists, tables, code blocks,
and formatting.
"""

import os
import re
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

BASE_DIR = Path(__file__).resolve().parent.parent
OUTPUT_DIR = BASE_DIR / "word"

# Palette
COLOR_PRIMARY = RGBColor(0x1F, 0x4E, 0x79)   # Navy
COLOR_SECONDARY = RGBColor(0x2E, 0x75, 0xB6) # Steel Blue
COLOR_DARK = RGBColor(0x26, 0x26, 0x26)      # Off-black
COLOR_MUTED = RGBColor(0x59, 0x59, 0x59)     # Gray


def set_cell_background(cell, hex_color):
    """Set background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set padding for a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def add_formatted_runs(paragraph, text, base_font="Times New Roman", base_size=Pt(11), base_color=COLOR_DARK):
    """Parse inline markdown (bold, italic, code) and add runs to paragraph."""
    # Pattern to match **bold**, *italic*, `code`
    token_pattern = re.compile(r'(\*\*[^*]+?\*\*|\*[^*]+?\*|`[^`]+?`)')
    parts = token_pattern.split(text)

    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            run.font.name = base_font
            run.font.size = base_size
            run.font.color.rgb = base_color
        elif part.startswith('*') and part.endswith('*') and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
            run.font.name = base_font
            run.font.size = base_size
            run.font.color.rgb = base_color
        elif part.startswith('`') and part.endswith('`') and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(base_size.pt * 0.92)
            run.font.color.rgb = RGBColor(0x99, 0x00, 0x00)
        else:
            run = paragraph.add_run(part)
            run.font.name = base_font
            run.font.size = base_size
            run.font.color.rgb = base_color


def process_table_lines(doc, table_lines):
    """Convert lines of Markdown table into a styled Word table."""
    if not table_lines:
        return

    # Parse rows
    parsed_rows = []
    for line in table_lines:
        # Check if delimiter line (e.g. |---|:---|)
        cleaned = line.strip().strip('|')
        cells = [c.strip() for c in cleaned.split('|')]
        # If it's a separator line like --- or :---: skip it
        if all(re.match(r'^:?-+:?$', c) for c in cells if c):
            continue
        parsed_rows.append(cells)

    if not parsed_rows:
        return

    num_cols = max(len(row) for row in parsed_rows)
    table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for r_idx, row in enumerate(parsed_rows):
        is_header = (r_idx == 0)
        for c_idx in range(num_cols):
            cell = table.cell(r_idx, c_idx)
            val = row[c_idx] if c_idx < len(row) else ""
            cell.text = ""  # clear
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05

            if is_header:
                set_cell_background(cell, "1F4E79")
                set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
                add_formatted_runs(p, val, base_font="Times New Roman", base_size=Pt(10), base_color=RGBColor(0xFF, 0xFF, 0xFF))
                for run in p.runs:
                    run.bold = True
            else:
                bg = "F2F5F8" if (r_idx % 2 == 1) else "FFFFFF"
                set_cell_background(cell, bg)
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                add_formatted_runs(p, val, base_font="Times New Roman", base_size=Pt(9.5), base_color=COLOR_DARK)

    # Empty paragraph after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(4)
    sp.paragraph_format.space_after = Pt(6)


def convert_md_to_docx(md_path: Path, output_docx_path: Path):
    """Convert a single Markdown file to a styled .docx file."""
    print(f"Converting: {md_path.name} -> {output_docx_path.name}")
    content = md_path.read_text(encoding="utf-8", errors="replace")
    lines = content.splitlines()

    doc = Document()

    # Document margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Normal style font
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    font.color.rgb = COLOR_DARK

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith("```"):
            if in_code_block:
                # End code block
                in_code_block = False
                code_text = "\n".join(code_lines)
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                set_cell_background(cell, "F4F5F7")
                set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.05
                run = p.add_run(code_text)
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                code_lines = []
                sp = doc.add_paragraph()
                sp.paragraph_format.space_before = Pt(2)
                sp.paragraph_format.space_after = Pt(4)
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Handle tables
        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_lines.append(stripped)
            i += 1
            continue
        elif in_table:
            # End of table
            in_table = False
            process_table_lines(doc, table_lines)
            table_lines = []

        # Empty lines
        if not stripped:
            i += 1
            continue

        # Horizontal rule
        if re.match(r'^(---+|\*\*\*+|___+)$', stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("―" * 55)
            run.font.color.rgb = RGBColor(0xBD, 0xC3, 0xC7)
            run.font.size = Pt(8)
            i += 1
            continue

        # Headings
        if stripped.startswith("#"):
            match = re.match(r'^(#{1,6})\s+(.*)$', stripped)
            if match:
                level = len(match.group(1))
                htext = match.group(2).strip()

                if level == 1:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(16)
                    p.paragraph_format.space_after = Pt(6)
                    p.paragraph_format.keep_with_next = True
                    add_formatted_runs(p, htext, base_font="Times New Roman", base_size=Pt(18), base_color=COLOR_PRIMARY)
                    for r in p.runs:
                        r.bold = True
                elif level == 2:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(12)
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.keep_with_next = True
                    add_formatted_runs(p, htext, base_font="Times New Roman", base_size=Pt(14), base_color=COLOR_SECONDARY)
                    for r in p.runs:
                        r.bold = True
                elif level == 3:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(10)
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.keep_with_next = True
                    add_formatted_runs(p, htext, base_font="Times New Roman", base_size=Pt(12), base_color=COLOR_PRIMARY)
                    for r in p.runs:
                        r.bold = True
                else:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(8)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.keep_with_next = True
                    add_formatted_runs(p, htext, base_font="Times New Roman", base_size=Pt(11), base_color=COLOR_DARK)
                    for r in p.runs:
                        r.bold = True
                i += 1
                continue

        # Blockquotes
        if stripped.startswith(">"):
            btext = re.sub(r'^>\s*', '', stripped)
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            set_cell_background(cell, "EBF3FA")
            set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_runs(p, btext, base_font="Times New Roman", base_size=Pt(10), base_color=COLOR_PRIMARY)
            for r in p.runs:
                r.italic = True
            i += 1
            continue

        # Bullet list items
        if re.match(r'^[-*+]\s+(.*)$', stripped):
            item_text = re.sub(r'^[-*+]\s+', '', stripped)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_runs(p, item_text, base_font="Times New Roman", base_size=Pt(11), base_color=COLOR_DARK)
            i += 1
            continue

        # Numbered list items
        num_match = re.match(r'^\d+\.\s+(.*)$', stripped)
        if num_match:
            item_text = num_match.group(1)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_runs(p, item_text, base_font="Times New Roman", base_size=Pt(11), base_color=COLOR_DARK)
            i += 1
            continue

        # Normal text paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        add_formatted_runs(p, stripped, base_font="Times New Roman", base_size=Pt(11), base_color=COLOR_DARK)
        i += 1

    # Cleanup table if file ended inside table
    if in_table and table_lines:
        process_table_lines(doc, table_lines)

    doc.save(str(output_docx_path))
    print(f"  -> Successfully created: {output_docx_path} ({os.path.getsize(output_docx_path):,} bytes)")


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    print(f"Output directory initialized at: {OUTPUT_DIR}")

    # Find all .md files in the repository, excluding .git and .pytest_cache
    md_files = []
    for root, dirs, files in os.walk(BASE_DIR):
        if ".git" in root or ".pytest_cache" in root or "node_modules" in root:
            continue
        for file in files:
            if file.endswith(".md"):
                md_files.append(Path(root) / file)

    print(f"Found {len(md_files)} Markdown files to convert:\n")

    converted_count = 0
    for md_file in sorted(md_files):
        rel_path = md_file.relative_to(BASE_DIR)

        # Name mapping to prevent collisions
        if rel_path == Path("README.md"):
            out_name = "README.docx"
        elif "data" in rel_path.parts:
            # e.g. data/raw/README.md -> data_raw_README.docx
            parent_parts = "_".join(rel_path.parent.parts)
            out_name = f"{parent_parts}_{md_file.stem}.docx"
        else:
            out_name = f"{md_file.stem}.docx"

        out_path = OUTPUT_DIR / out_name
        try:
            convert_md_to_docx(md_file, out_path)
            converted_count += 1
        except Exception as e:
            print(f"  [ERROR] Failed to convert {md_file}: {e}")

    print("\n" + "=" * 70)
    print(f"ALL CONVERSIONS COMPLETE: {converted_count}/{len(md_files)} files converted successfully.")
    print(f"All .docx files are saved in: {OUTPUT_DIR}")
    print("=" * 70)


if __name__ == "__main__":
    main()
